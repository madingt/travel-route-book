#!/usr/bin/env python3
"""Export trip.json to a spreadsheet: one row per place + a per-day summary with live formulas.

Usage: python3 trip_to_xlsx.py trip.json out.xlsx
Needs openpyxl. Counts on the summary sheet are COUNTIF/COUNTIFS formulas, so they update
when the place sheet is edited (recalculate in Excel/Numbers/LibreOffice on open).
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from trip_lib import load_trip, place_url  # noqa: E402

FONT = "Arial"
PALETTE = ["EAF4F4", "FFF4E5", "EEF0FB", "EAF6EC", "FBEFF3", "F2F2F2", "FFF9DB"]
TYPES = ["景点", "餐饮", "咖啡", "交通", "住宿", "活动"]


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    trip = load_trip(sys.argv[1])
    days = trip["days"]
    map_label = "谷歌地图" if trip.get("map") == "google" else "高德"

    hdr_font = Font(name=FONT, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", start_color="2F5D62")
    body = Font(name=FONT)
    link = Font(name=FONT, color="0563C1", underline="single")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = {d["id"]: PatternFill("solid", start_color=PALETTE[i % len(PALETTE)]) for i, d in enumerate(days)}

    wb = Workbook()
    ws = wb.active
    ws.title = "地点"
    headers = ["天", "日期", "分组名", "组内序号", "总序号", "时间", "地点", "类型", "安排/备注",
               "地址", "营业时间", "状态", "地图链接"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    r = 2
    for d in days:
        group = f'{d.get("label", d["id"])}·{d.get("date", "")}·{d.get("title", "")}'.strip("·")
        for n, s in enumerate(d.get("stops", []), start=1):
            vals = [d.get("label", d["id"]), d.get("date", ""), group, n, "=ROW()-1", s.get("time", ""),
                    s["name"], s.get("type", ""), s.get("note", ""), s.get("address", ""),
                    s.get("hours", ""), s.get("status", "")]
            for c, v in enumerate(vals, start=1):
                ws.cell(row=r, column=c, value=v)
            url = place_url(trip, s).replace('"', '""')
            ws.cell(row=r, column=13, value=f'=HYPERLINK("{url}","在{map_label}打开")')
            for c in range(1, 14):
                cell = ws.cell(row=r, column=c)
                cell.font = link if c == 13 else body
                cell.fill, cell.border = fill[d["id"]], border
                cell.alignment = Alignment(vertical="center", wrap_text=c in (3, 7, 9, 10, 11),
                                           horizontal="center" if c in (1, 4, 5, 6, 8, 12) else "left")
            r += 1
    last = r - 1
    for c, w in enumerate([6, 12, 26, 9, 8, 8, 30, 8, 26, 36, 18, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(last, 1)}"
    ws.cell(row=last + 2, column=1, value=f"来源：{trip['title']} 的 trip.json；营业时间与预约以官方渠道为准。"
            ).font = Font(name=FONT, italic=True, color="666666")

    s = wb.create_sheet("每日概览")
    s.append(["天", "日期", "主题", "地点数"] + TYPES)
    for c in range(1, 5 + len(TYPES)):
        cell = s.cell(row=1, column=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(horizontal="center")
    rng_d, rng_t = f"'地点'!$A$2:$A${last}", f"'地点'!$H$2:$H${last}"
    for i, d in enumerate(days, start=2):
        s.cell(row=i, column=1, value=d.get("label", d["id"]))
        s.cell(row=i, column=2, value=d.get("date", ""))
        s.cell(row=i, column=3, value=d.get("title", ""))
        s.cell(row=i, column=4, value=f"=COUNTIF({rng_d},A{i})")
        for j, t in enumerate(TYPES, start=5):
            s.cell(row=i, column=j, value=f'=COUNTIFS({rng_d},$A{i},{rng_t},"{t}")')
    tr = len(days) + 2
    s.cell(row=tr, column=1, value="合计")
    for c in range(4, 5 + len(TYPES)):
        col = get_column_letter(c)
        s.cell(row=tr, column=c, value=f"=SUM({col}2:{col}{tr - 1})")
    for row in range(2, tr + 1):
        for c in range(1, 5 + len(TYPES)):
            cell = s.cell(row=row, column=c)
            cell.font, cell.border = Font(name=FONT, bold=(row == tr)), border
            cell.alignment = Alignment(horizontal="left" if c == 3 else "center")
            if row < tr:
                cell.fill = fill[days[row - 2]["id"]]
    for c, w in enumerate([6, 12, 30, 8] + [8] * len(TYPES), start=1):
        s.column_dimensions[get_column_letter(c)].width = w

    wb.save(sys.argv[2])
    print(f"wrote {sys.argv[2]} ({last - 1} places, {len(days)} days)")


if __name__ == "__main__":
    main()
